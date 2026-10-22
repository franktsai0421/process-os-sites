"""Audit the Local IQC workbook and optionally build a row-anchored source manifest.

Read-only by default. The workbook and any existing Local database are never edited.
"""

import argparse
import hashlib
import json
import re
import sys
from collections import Counter, defaultdict
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
DEFAULT_SOURCE = ROOT / "KD" / "IQC" / "IQC_CHUAN 的複本.xlsx"
DEFAULT_MAPPING = ROOT / "KD" / "IQC" / "iqc_category_mapping.json"
OUTPUT = ROOT / "public" / "process" / "iqc-source.json"
PHOTO_DIR = ROOT / "public" / "process" / "iqc-media"


def source_text(value):
    return "" if value is None else str(value)


def cell_values(sheet):
    merged = {}
    for area in sheet.merged_cells.ranges:
        top_left = sheet.cell(area.min_row, area.min_col).value
        for row in range(area.min_row, area.max_row + 1):
            for col in range(area.min_col, area.max_col + 1):
                merged[row, col] = top_left

    def get(row, col):
        value = sheet.cell(row, col).value
        return source_text(value if value is not None else merged.get((row, col)))

    return get


def read_source(path):
    workbook = load_workbook(path, data_only=True)
    sheet = workbook["IQC_MASTER"]
    get = cell_values(sheet)
    parts = {}
    errors = []
    items = []
    issues = []
    current = None
    parent = {}
    row_to_code = {}
    for row in range(2, sheet.max_row + 1):
        code = get(row, 1).strip()
        error_code = get(row, 8).strip()
        check_no = get(row, 10).strip()
        if not code and not error_code and not check_no:
            continue
        if code:
            if code != current:
                if code in parts:
                    issues.append({"row": row, "type": "NONCONTIGUOUS_PART", "partCode": code})
                current = code
                parent = {}
            for col in (2, 3, 5, 6, 7, 15, 16, 17, 18, 19, 20):
                value = get(row, col)
                if value:
                    parent[col] = value
        if not current:
            issues.append({"row": row, "type": "ORPHAN_ROW"})
            continue
        row_to_code[row] = current
        if current not in parts:
            parts[current] = {
                "id": f"iqc-part:{current}", "code": current, "nameVi": parent.get(2, ""),
                "specificationVi": parent.get(3, ""), "inspectionType": parent.get(5, ""),
                "supplierRaw": parent.get(6, ""), "standardCode": parent.get(7, ""),
                "sourceRow": row, "sourceRows": [], "photos": [], "categoryId": "",
            }
        part = parts[current]
        part["sourceRows"].append(row)
        for col, key in ((2, "nameVi"), (3, "specificationVi"), (5, "inspectionType"), (6, "supplierRaw"), (7, "standardCode")):
            value = get(row, col) or parent.get(col, "")
            if value and part[key] and value != part[key]:
                issues.append({"row": row, "type": "PARENT_VALUE_CHANGED", "partCode": current, "field": key, "value": value})
            elif value:
                part[key] = value
        if not error_code or not check_no:
            issues.append({"row": row, "type": "MISSING_ERROR_OR_CHECK", "partCode": current})
            continue
        error = {"id": f"iqc-error:{error_code}", "partCode": current, "code": error_code,
                 "nameVi": get(row, 9), "sourceRow": row, "isDeleted": False}
        item = {"id": f"iqc-item:{check_no}", "partCode": current, "errorId": error["id"],
                "checkNo": check_no, "nameVi": get(row, 11), "toolVi": get(row, 12),
                "methodVi": get(row, 13), "frequencyRaw": get(row, 14),
                "enabledRaw": get(row, 15), "versionRaw": get(row, 16),
                "dateRaw": get(row, 17), "reasonVi": get(row, 18),
                "creatorVi": get(row, 19), "noteVi": get(row, 20),
                "sourceRow": row, "isDeleted": False}
        errors.append(error)
        items.append(item)
        if not item["nameVi"].strip():
            issues.append({"row": row, "type": "BLANK_CHECK_NAME", "partCode": current,
                           "message": f"原始檢查項目 {check_no} 名稱空白；未自動補寫"})

    image_anchors = []
    for image in sheet._images:
        anchor = image.anchor._from
        row, col = anchor.row + 1, anchor.col + 1
        code = row_to_code.get(row)
        raw = image._data()
        digest = hashlib.sha256(raw).hexdigest()
        image_anchors.append({"row": row, "column": col, "partCode": code,
                              "sha256": digest, "bytes": len(raw), "format": image.format})
        if not code or col != 4:
            issues.append({"row": row, "type": "IMAGE_ANCHOR_UNMATCHED", "column": col})
            continue
        parts[code]["photos"].append({"sourceRow": row, "sourceColumn": col,
                                      "sha256": digest, "format": image.format})

    error_codes = [item["code"] for item in errors]
    check_nos = [item["checkNo"] for item in items]
    for kind, values in (("DUPLICATE_ERROR_CODE", error_codes), ("DUPLICATE_CHECK_NO", check_nos)):
        for code, count in Counter(values).items():
            if count > 1:
                issues.append({"type": kind, "code": code, "count": count})
    for part in parts.values():
        if not part["photos"]:
            issues.append({"row": part["sourceRow"], "type": "NO_SOURCE_IMAGE", "partCode": part["code"]})

    suppliers = Counter(part["supplierRaw"] for part in parts.values())
    summary = {"source": str(path.relative_to(ROOT)), "sha256": hashlib.sha256(path.read_bytes()).hexdigest(),
               "rows": len(items), "parts": len(parts), "errors": len(errors),
               "uniqueErrorCodes": len(set(error_codes)), "uniqueCheckNos": len(set(check_nos)),
               "images": len(image_anchors), "partsWithImages": sum(bool(p["photos"]) for p in parts.values()),
               "suppliers": suppliers, "issues": issues,
               "imageAnchors": image_anchors}
    return {"summary": summary, "parts": list(parts.values()), "errors": errors, "items": items}


def supplier_names(raw):
    if not raw:
        return []
    # The workbook uses commas for explicit lists, including one three-supplier part.
    # A plus sign or name-internal hyphen is not a safe split rule.
    return [name.strip() for name in raw.split(",") if name.strip()]


def category_mapping(path, parts):
    if not path.is_file():
        return [], {}, [{"type": "CATEGORY_MAPPING_MISSING", "message": "iqc_category_mapping.json 不存在；尚未指定任何零件類別"}], ""
    raw = path.read_bytes()
    mapping = json.loads(raw.decode("utf-8-sig"))
    if not isinstance(mapping, dict):
        raise ValueError("類別對照 JSON 頂層必須是物件")
    category_rows = mapping.get("categories")
    part_rows = mapping.get("parts", mapping.get("mappings"))
    if isinstance(category_rows, dict):
        category_rows = [{"id": key, "nameVi": value} for key, value in category_rows.items()]
    if not isinstance(category_rows, list) or not isinstance(part_rows, (list, dict)):
        raise ValueError("類別對照需要 categories 與 parts/mappings；請核對實際 JSON 格式")
    categories = []
    for row in category_rows:
        if isinstance(row, str):
            row = {"nameVi": row}
        name = source_text(row.get("nameVi", row.get("name_vi", row.get("name", "")))).strip()
        ident = source_text(row.get("id", "")).strip() or f"iqc-category:{hashlib.sha256(name.encode()).hexdigest()[:12]}"
        if not name or any(item["id"] == ident or item["nameVi"] == name for item in categories):
            raise ValueError("類別對照有空白或重複的八類名稱／ID")
        categories.append({"id": ident, "nameVi": name})
    if isinstance(part_rows, dict):
        part_rows = [{"code": code, **(value if isinstance(value, dict) else {"categoryId": value})} for code, value in part_rows.items()]
    by_code = {}
    for row in part_rows:
        code = source_text(row.get("code", row.get("partCode", row.get("part_code", "")))).strip()
        if not code or code in by_code:
            raise ValueError("類別對照零件編號空白或重複")
        by_code[code] = row
    category_ids = {item["id"] for item in categories}
    category_names = {item["nameVi"]: item["id"] for item in categories}
    assignments = {}
    issues = []
    for part in parts:
        row = by_code.get(part["code"])
        if row is None:
            issues.append({"row": part["sourceRow"], "partCode": part["code"], "type": "CATEGORY_MAPPING_MISSING_PART", "message": "類別對照缺少此零件"})
            continue
        mapped_name = source_text(row.get("nameVi", row.get("partNameVi", row.get("source_name_vi", row.get("name", ""))))).strip()
        # The reviewed mapping uses single spaces where Excel wraps a name onto
        # another visual line. Compare the text, not its cell line layout.
        if re.sub(r"\s+", " ", mapped_name) != re.sub(r"\s+", " ", part["nameVi"].strip()):
            issues.append({"row": part["sourceRow"], "partCode": part["code"], "type": "CATEGORY_NAME_MISMATCH", "message": f"Excel：{part['nameVi']}；JSON：{mapped_name}"})
            continue
        requested = source_text(row.get("categoryId", row.get("category_id", row.get("category", row.get("categoryNameVi", ""))))).strip()
        category_id = requested if requested in category_ids else category_names.get(requested, "")
        if not category_id:
            issues.append({"row": part["sourceRow"], "partCode": part["code"], "type": "CATEGORY_UNKNOWN", "message": f"類別對照找不到八類中的：{requested}"})
            continue
        assignments[part["code"]] = category_id
    for code in by_code.keys() - {part["code"] for part in parts}:
        issues.append({"partCode": code, "type": "CATEGORY_EXTRA_PART", "message": "JSON 零件編號不在 Excel"})
    if len(categories) != 8:
        issues.append({"type": "CATEGORY_COUNT", "message": f"對照檔只有 {len(categories)} 類，預期八類"})
    return categories, assignments, issues, hashlib.sha256(raw).hexdigest()


def build_manifest(source_path, mapping_path, write=False):
    extracted = read_source(source_path)
    parts = extracted["parts"]
    categories, assignments, mapping_issues, mapping_sha = category_mapping(mapping_path, parts)
    issues = list(extracted["summary"]["issues"]) + mapping_issues
    if write:
        PHOTO_DIR.mkdir(parents=True, exist_ok=True)
    workbook = load_workbook(source_path, data_only=True)
    sheet = workbook["IQC_MASTER"]
    photos_by_row = defaultdict(list)
    for image in sheet._images:
        row, col = image.anchor._from.row + 1, image.anchor._from.col + 1
        if col != 4:
            raise ValueError(f"照片第 {row} 列錨點不在第 4 欄")
        raw = image._data()
        digest = hashlib.sha256(raw).hexdigest()
        relative = f"/process/iqc-media/row-{row}-{digest[:12]}.jpg"
        if write:
            path = ROOT / "public" / relative.lstrip("/")
            if path.exists() and hashlib.sha256(path.read_bytes()).hexdigest() != digest:
                raise ValueError(f"已有照片與來源不符：{path}")
            if not path.exists():
                path.write_bytes(raw)
        photos_by_row[row].append({"id": digest, "src": relative, "webPath": relative,
                                   "sourceRow": row, "sourceColumn": col, "name": f"Excel row {row}"})
    suppliers = {}
    snapshot_parts = []
    for part in parts:
        raw = part["supplierRaw"]
        names = supplier_names(raw)
        if " + " in raw or raw == "HƯNG LONG-TÂN NGỌC LAN":
            names = []
            issues.append({"row": part["sourceRow"], "partCode": part["code"], "type": "SUPPLIER_SPLIT_UNCERTAIN", "message": f"保留複合供應商原文：{raw}"})
        if part["code"] == "3080601XI":
            # The request says this supplier is blank, while the supplied workbook is not.
            # Keep the actual cell text, but leave the structured assignment for review.
            names = []
            issues.append({"row": part["sourceRow"], "partCode": part["code"], "type": "SUPPLIER_SOURCE_CONFLICT", "message": f"文字需求稱原始供應商空白；Excel F{part['sourceRow']} 實為：{raw}。未自動指定供應商"})
        supplier_ids = []
        for name in names:
            ident = f"iqc-supplier:{hashlib.sha256(name.encode()).hexdigest()[:12]}"
            suppliers[ident] = {"id": ident, "nameVi": name}
            supplier_ids.append(ident)
        snapshot_parts.append({"id": part["id"], "code": part["code"], "nameVi": part["nameVi"],
                               "specificationVi": part["specificationVi"],
                               "categoryId": assignments.get(part["code"], ""),
                               "supplierIds": supplier_ids, "supplierRaw": raw,
                               "standardCode": part["standardCode"], "inspectionType": part["inspectionType"],
                               "sourceRow": part["sourceRow"], "sourceRows": part["sourceRows"],
                               "photos": photos_by_row.get(part["sourceRow"], [])})
    errors = extracted["errors"]
    items = []
    for item in extracted["items"]:
        frequency = item["frequencyRaw"]
        items.append({"id": item["id"], "partCode": item["partCode"], "errorId": item["errorId"],
                      "checkNo": item["checkNo"], "nameVi": item["nameVi"],
                      "toolVi": item["toolVi"], "methodVi": item["methodVi"],
                      "frequencyRaw": frequency, "sourceRow": item["sourceRow"],
                      "isEnabled": item["enabledRaw"].strip().upper() == "YES",
                      "source": {"enabledRaw": item["enabledRaw"], "versionRaw": item["versionRaw"],
                                 "dateRaw": item["dateRaw"], "reasonVi": item["reasonVi"],
                                 "creatorVi": item["creatorVi"], "noteVi": item["noteVi"]},
                      "rules": [{"id": f"iqc-source-rule:{item['checkNo']}", "cadence": "source",
                                 "kind": "source", "value": "", "sourceTextVi": frequency}]})
    for item in ("NƯỚC NGOÀI", "CÔNG TY"):
        ident = f"iqc-supplier:{hashlib.sha256(item.encode()).hexdigest()[:12]}"
        suppliers.setdefault(ident, {"id": ident, "nameVi": item})
    manifest = {"ready": bool(mapping_sha and len(categories) == 8 and assignments), "fileName": source_path.name,
                "sha256": extracted["summary"]["sha256"], "mappingSha256": mapping_sha,
                "snapshot": {"categories": categories, "suppliers": list(suppliers.values()),
                             "parts": snapshot_parts, "errors": errors, "items": items},
                "issues": issues}
    if write:
        OUTPUT.write_text(json.dumps(manifest, ensure_ascii=False, separators=(",", ":")), encoding="utf-8")
    return manifest


def main():
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")
    parser = argparse.ArgumentParser()
    parser.add_argument("--source", type=Path, default=DEFAULT_SOURCE)
    parser.add_argument("--mapping", type=Path, default=DEFAULT_MAPPING)
    parser.add_argument("--summary", action="store_true", help="print compact read-only source audit")
    parser.add_argument("--write", action="store_true", help="write local preview manifest and row-anchored JPEG files")
    args = parser.parse_args()
    result = read_source(args.source)
    summary = result["summary"]
    if args.summary:
        compact = {key: value for key, value in summary.items() if key != "imageAnchors"}
        compact["noImageParts"] = [issue["partCode"] for issue in summary["issues"] if issue["type"] == "NO_SOURCE_IMAGE"]
        print(json.dumps(compact, ensure_ascii=False, indent=2, default=dict))
    if args.write:
        manifest = build_manifest(args.source, args.mapping, write=True)
        print(json.dumps({"ready": manifest["ready"], "parts": len(manifest["snapshot"]["parts"]),
                          "categories": len(manifest["snapshot"]["categories"]),
                          "errors": len(manifest["snapshot"]["errors"]),
                          "items": len(manifest["snapshot"]["items"]),
                          "photos": sum(len(p["photos"]) for p in manifest["snapshot"]["parts"]),
                          "issues": len(manifest["issues"])}, ensure_ascii=False))


if __name__ == "__main__":
    main()
