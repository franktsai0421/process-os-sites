"""Pinned workbook, read-only audit against explicitly captured effective Local AB master."""
import argparse
import hashlib
import json
import re
import unicodedata
from collections import Counter, defaultdict
from pathlib import Path
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
SHA = "e5bb57abfbe5ea9221a7e15ed1fd8ab964b9211968fdae872facc7ece8c712b3"
SOURCE = ROOT / "KD/IPQC-AB/01_最新資料/IPQC_CHUAN_29092026_V.xlsx"
FIELDS = {"methodText": 10, "sourceTool": 12, "sourceMethod": 13, "sourceFrequency": 14,
          "sourceEnabled": 15, "sourceVersion": 16, "sourceChangedAt": 17,
          "sourceReason": 18, "sourceCreator": 19, "sourceNote": 20}

def text(cell):
    if cell.value is None: return ""
    if cell.data_type == "f": raise ValueError(f"Formula in import field {cell.coordinate}")
    value = cell.value
    if hasattr(value, "isoformat"): return value.isoformat()
    if isinstance(value, (float, int)) and not isinstance(value, bool):
        value = str(int(value)) if int(value) == value else str(value)
        if re.fullmatch(r"0+", cell.number_format): value = value.zfill(len(cell.number_format))
    return str(value).replace("\r\n", "\n").replace("\r", "\n").strip()

def main():
    cli = argparse.ArgumentParser()
    cli.add_argument("--source", type=Path, default=SOURCE)
    cli.add_argument("--catalog", type=Path, required=True)
    cli.add_argument("--report", type=Path, required=True)
    args = cli.parse_args()
    if args.source.resolve() != SOURCE.resolve(): raise ValueError("Only latest authorized workbook is accepted")
    digest = hashlib.sha256(args.source.read_bytes()).hexdigest()
    if digest != SHA: raise ValueError("Source SHA256 differs; stopped")
    catalog = json.loads(args.catalog.read_text(encoding="utf-8-sig"))
    groups = {(g["line"], g["name"]) for g in catalog["groups"]}
    stages = {(s["line"], s["code"]) for s in catalog["stages"]}
    products = defaultdict(list)
    for product in catalog["products"]: products[product["code"]].append(product)
    sheet = load_workbook(args.source, data_only=False)["IPQC"]
    merged, anchors = {}, {}
    for region in sheet.merged_cells.ranges:
        value = text(sheet.cell(region.min_row, region.min_col))
        for row in range(region.min_row, region.max_row + 1):
            for col in range(region.min_col, min(20, region.max_col) + 1):
                merged[row, col] = value
                if col == 8: anchors[row] = str(region)
    def cell(row, col): return merged.get((row, col), text(sheet.cell(row, col)))
    rows, variants = [], {}
    for row in range(2, sheet.max_row + 1):
        values = [cell(row, col) for col in range(1, 9)]
        if not any(values): continue
        line_raw, group, stage, error, name, item, standard, raw_products = values
        line = line_raw.removeprefix("LINE ").strip()
        basic = {key: cell(row, col) for key, col in FIELDS.items()}
        identity = [line, group, f"{line}-{stage}", error, unicodedata.normalize("NFC", standard), *basic.values()]
        key = hashlib.sha256(json.dumps(identity, ensure_ascii=False, separators=(",", ":")).encode()).hexdigest()
        variant = variants.setdefault(key, dict(key=key, line=line, group=group, stageCode=f"{line}-{stage}",
            errorCode=error, errorName=name, standard=standard, **basic, sourceRows=[], sourceItemCodes=[], productCodes=[]))
        variant["sourceRows"].append(row)
        variant["sourceItemCodes"].append(item)
        for raw in raw_products.splitlines() or [""]:
            raw = raw.strip()
            code = raw.split(maxsplit=1)[0] if raw else ""
            reasons = []
            if line not in ("A", "B"): reasons.append("line_missing")
            if (line, group) not in groups: reasons.append("group_unmapped")
            if (line, variant["stageCode"]) not in stages: reasons.append("stage_missing")
            matches = products.get(code, [])
            if not matches: reasons.append("product_missing")
            elif len(matches) != 1: reasons.append("product_ambiguous")
            else:
                product = matches[0]
                if product["line"] != line: reasons.append("product_line_mismatch")
                if product["category"] != group: reasons.append("product_group_mismatch")
                if variant["stageCode"] not in product["stages"]: reasons.append("product_route_mismatch")
            if not re.fullmatch(re.escape(f"{line}-{group}-{stage}-L") + r"\d+", error): reasons.append("error_context_mismatch")
            if not re.fullmatch(re.escape(error + "-") + r"\d+", item): reasons.append("code_parent_mismatch")
            if basic["sourceEnabled"] not in ("YES", "NO"): reasons.append("enabled_invalid")
            if not code: reasons.append("product_code_blank")
            rows.append(dict(row=row, key=key, line=line, sourceGroup=group, stageCode=variant["stageCode"],
                errorCode=error, errorName=name, itemCode=item, standard=standard, productCode=code,
                sourceProduct=raw, productAnchor=anchors.get(row, f"H{row}"), **basic, reasons=reasons))
    standard_groups, basics, errors = defaultdict(list), defaultdict(list), defaultdict(set)
    for variant in variants.values():
        for field in ("sourceRows", "sourceItemCodes"): variant[field] = sorted(set(variant[field]))
        basics[(variant["errorCode"], unicodedata.normalize("NFC", variant["standard"]))].append(variant)
    for row in rows:
        standard_groups[row["itemCode"]].append(row)
        errors[row["errorCode"]].add((row["line"], row["sourceGroup"], row["stageCode"], row["errorName"]))
    basic_conflicts = [dict(errorCode=code, standard=standard, keys=[v["key"] for v in values],
        differingFields=[f for f in FIELDS if len({v[f] for v in values}) > 1],
        sourceRows=sorted({r for v in values for r in v["sourceRows"]}))
        for (code, standard), values in basics.items() if len(values) > 1]
    error_conflicts = [dict(errorCode=code, contexts=sorted(contexts)) for code, contexts in errors.items() if len(contexts) > 1]
    blocked_keys = {key for c in basic_conflicts for key in c["keys"]}
    blocked_errors = {c["errorCode"] for c in error_conflicts}
    by_key = defaultdict(list)
    for row in rows:
        if row["key"] in blocked_keys: row["reasons"].append("shared_basic_fields_conflict")
        if row["errorCode"] in blocked_errors: row["reasons"].append("error_context_conflict")
        by_key[row["key"]].append(row)
    for key, variant in variants.items():
        variant["productCodes"] = sorted({r["productCode"] for r in by_key[key] if not r["reasons"]})
        variant["blocked"] = key in blocked_keys or variant["errorCode"] in blocked_errors
    conflicts = [dict(itemCode=code, variants=[dict(key=key, standard=variants[key]["standard"],
                  sourceRows=variants[key]["sourceRows"], productCodes=variants[key]["productCodes"]) for key in sorted({r["key"] for r in entries})])
                 for code, entries in standard_groups.items() if len({r["standard"] for r in entries}) > 1]
    accepted = [v for v in variants.values() if v["productCodes"]]
    valid = sum(not r["reasons"] for r in rows)
    unique = sum(len(v["productCodes"]) for v in accepted)
    summary = dict(physicalRows=len({r["row"] for r in rows}), sourceGroups=len({(r["line"],r["sourceGroup"]) for r in rows}),
        sourceErrors=len(errors), sourceItemNumbers=len(standard_groups), standardVariants=len(variants),
        sourceProducts=len({r["productCode"] for r in rows}), expandedRelations=len(rows), acceptedExpandedRelations=valid,
        excludedExpandedRelations=len(rows)-valid, uniqueValidRelations=unique, duplicateExpandedRelations=valid-unique,
        acceptedErrors=len({v["errorCode"] for v in accepted}), acceptedItems=len(accepted),
        itemsWithoutValidProducts=len(variants)-len(accepted), standardConflictNumbers=len(conflicts),
        basicConflicts=len(basic_conflicts), errorConflicts=len(error_conflicts),
        reasonCounts=dict(Counter(reason for r in rows for reason in r["reasons"])))
    unresolved = [dict(line=line, sourceGroup=group, sourceRows=sorted({r["row"] for r in rows if r["line"] == line and r["sourceGroup"] == group}))
                  for line, group in sorted({(r["line"],r["sourceGroup"]) for r in rows} - groups)]
    report = dict(format="ipqc-ab-audit/v2", source=dict(path=str(args.source.resolve()), sha256=digest, fileName=args.source.name,
        sheet="IPQC", rows=sheet.max_row-1, mergedRanges=len(sheet.merged_cells.ranges), embeddedImages=len(sheet._images)),
        catalog=catalog, summary=summary, rows=rows, variants=list(variants.values()), standardConflicts=conflicts,
        basicConflicts=basic_conflicts, errorConflicts=error_conflicts, unresolvedGroups=unresolved)
    args.report.parent.mkdir(parents=True, exist_ok=True)
    args.report.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(summary, ensure_ascii=False, indent=2))

if __name__ == "__main__": main()
