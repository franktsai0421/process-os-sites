"""Derive the reviewed Local import fixture from the original equipment workbook.

The workbook remains the source of truth. Images are matched by their Excel
drawing anchors (row and photo column), never by their ZIP/media order.
"""
import hashlib
import io
import json
from collections import Counter
from datetime import date, datetime
from pathlib import Path

import openpyxl
from PIL import Image


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / 'KD/ROUTINE/原始設備Excel.xlsx'
OUT = ROOT / 'public/process/routine-source.json'
MEDIA = ROOT / 'public/process/routine-media'


def value(cell):
    raw = cell.value
    return raw.isoformat()[:10] if isinstance(raw, (date, datetime)) else raw or ''


def main():
    book = openpyxl.load_workbook(SOURCE, data_only=True)
    master, _, _, group_sheet, _, issue_sheet, proof_sheet = book.worksheets[:7]
    groups = {}
    for row in range(2, group_sheet.max_row + 1):
        key = value(group_sheet.cell(row, 1))
        if key:
            groups[key] = {
                'id': key, 'nameZh': value(group_sheet.cell(row, 2)),
                'nameVi': value(group_sheet.cell(row, 3)),
                'categoryCode': value(group_sheet.cell(row, 4)),
                'category': value(group_sheet.cell(row, 5)),
                'subcategory': value(group_sheet.cell(row, 6)),
                'spec': value(group_sheet.cell(row, 7)),
                'sourceStatus': value(group_sheet.cell(row, 8)),
                'notes': value(group_sheet.cell(row, 9)),
            }
    images = {}
    for image in master._images:
        row, column = image.anchor._from.row + 1, image.anchor._from.col + 1
        if column != 18 or row in images:
            raise ValueError(f'Unexpected or duplicate photo anchor: {row}, {column}')
        images[row] = image
    MEDIA.mkdir(parents=True, exist_ok=True)
    equipment = []
    for row in range(2, 121):
        asset = value(master.cell(row, 4))
        if not asset or row not in images:
            raise ValueError(f'Missing asset or photo at row {row}')
        original = images[row]._data()
        output = io.BytesIO()
        Image.open(io.BytesIO(original)).convert('RGB').save(output, format='WEBP', quality=78, method=6)
        filename = f'row-{row}.webp'
        (MEDIA / filename).write_bytes(output.getvalue())
        proof = [value(proof_sheet.cell(row, column)) for column in range(1, 7)]
        if proof[0] != asset:
            raise ValueError(f'Proof/asset mismatch at row {row}')
        equipment.append({
            'id': f'xlsx-row-{row}', 'sourceRow': row, 'assetCode': asset,
            'pieceMark': value(master.cell(row, 23)),
            'categoryCode': value(master.cell(row, 1)),
            'category': value(master.cell(row, 2)),
            'subcategory': value(master.cell(row, 3)),
            'groupId': value(master.cell(row, 22)),
            'nameZh': value(master.cell(row, 5)),
            'nameVi': value(master.cell(row, 6)),
            'quantity': value(master.cell(row, 7)),
            'unit': value(master.cell(row, 8)),
            'spec': value(master.cell(row, 9)),
            'owner': value(master.cell(row, 10)),
            'sourceStatus': value(master.cell(row, 11)),
            'enabled': not str(value(master.cell(row, 11))).startswith('未啟用'),
            'calibrationMethod': value(master.cell(row, 12)),
            'calibrationCycle': value(master.cell(row, 13)),
            'calibratedAt': value(master.cell(row, 14)),
            'expiresAt': value(master.cell(row, 15)),
            'sourceHistory': value(master.cell(row, 16)),
            'sourceStandardLinks': value(master.cell(row, 17)),
            'notes': value(master.cell(row, 19)),
            'customCycle': value(master.cell(row, 20)),
            'intakeDate': value(master.cell(row, 21)),
            'sourceProductLinks': value(master.cell(row, 24)),
            'sourceReference': proof[3],
            'photoSource': {
                'workbookRow': row, 'column': 18,
                'originalSha256': hashlib.sha256(original).hexdigest(),
                'webPath': f'/process/routine-media/{filename}',
                'webSha256': hashlib.sha256(output.getvalue()).hexdigest(),
            },
            'confirmed': False,
        })
    issues = []
    for row in range(2, issue_sheet.max_row + 1):
        if not value(issue_sheet.cell(row, 1)):
            continue
        issues.append({
            'sourceRow': row,
            'assetCode': value(issue_sheet.cell(row, 1)),
            'pieceMark': value(issue_sheet.cell(row, 2)),
            'name': value(issue_sheet.cell(row, 3)),
            'type': value(issue_sheet.cell(row, 4)),
            'detail': value(issue_sheet.cell(row, 5)),
            'status': value(issue_sheet.cell(row, 6)),
        })
    count = Counter(item['assetCode'] for item in equipment)
    assert len(equipment) == 119 and len(images) == 119 and len(issues) == 37
    assert len([v for v in count.values() if v > 1]) == 9
    assert sum(v for v in count.values() if v > 1) == 30
    assert sum(not item['enabled'] for item in equipment) == 6
    manifest = {
        'schemaVersion': 1, 'fileName': SOURCE.name,
        'sha256': hashlib.sha256(SOURCE.read_bytes()).hexdigest(),
        'equipment': equipment, 'groups': list(groups.values()), 'issues': issues,
    }
    OUT.write_text(json.dumps(manifest, ensure_ascii=False, separators=(',', ':')), encoding='utf-8')
    print(json.dumps({'rows': len(equipment), 'photos': len(images), 'issues': len(issues),
                      'duplicateAssetGroups': 9, 'duplicateAssetRows': 30,
                      'sha256': manifest['sha256']}, ensure_ascii=False))


if __name__ == '__main__':
    main()
